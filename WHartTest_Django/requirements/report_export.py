# -*- coding: utf-8 -*-
"""评审报告导出：生成 Word(docx) 与 PDF 字节流。

- Word: python-docx
- PDF: reportlab(自带 STSong-Light CID 字体，无需外部中文字体文件)
"""
from __future__ import annotations

import io

from .models import ReviewReport


_RATING_FALLBACK = {
    "excellent": "优秀",
    "good": "良好",
    "average": "一般",
    "fair": "一般",
    "needs_improvement": "需改进",
    "poor": "较差",
}


def _rating_text(report: ReviewReport) -> str:
    try:
        return report.get_overall_rating_display()
    except Exception:
        return _RATING_FALLBACK.get(report.overall_rating, report.overall_rating or "-")


def _score_rows(report: ReviewReport):
    return [
        ("完整度", report.completion_score),
        ("清晰度", report.clarity_score),
        ("一致性", report.consistency_score),
        ("完整性", report.completeness_score),
        ("可测性", report.testability_score),
        ("可行性", report.feasibility_score),
        ("逻辑性", report.logic_score),
    ]


def _issue_display(issue, attr: str, default: str = "-") -> str:
    getter = getattr(issue, f"get_{attr}_display", None)
    if callable(getter):
        try:
            return getter()
        except Exception:
            pass
    return getattr(issue, attr, None) or default


def _report_filename_stem(report: ReviewReport) -> str:
    title = getattr(getattr(report, "document", None), "title", None) or "评审报告"
    date = report.review_date.strftime("%Y%m%d") if report.review_date else ""
    stem = f"{title}_评审报告_{date}".strip("_")
    # 去掉文件名非法字符
    for ch in '\\/:*?"<>|\n\r\t':
        stem = stem.replace(ch, "_")
    return stem or "review_report"


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def build_docx(report: ReviewReport) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc_title = getattr(getattr(report, "document", None), "title", None) or "需求评审报告"
    doc.add_heading(f"{doc_title} — 需求评审报告", level=0)

    meta = doc.add_paragraph()
    meta.add_run("评审人：").bold = True
    meta.add_run(report.reviewer or "AI需求评审助手")
    meta.add_run("    评审时间：").bold = True
    meta.add_run(report.review_date.strftime("%Y-%m-%d %H:%M") if report.review_date else "-")
    meta.add_run("    总体评价：").bold = True
    meta.add_run(_rating_text(report))

    # 评分
    doc.add_heading("评分概览", level=1)
    score_table = doc.add_table(rows=1, cols=2)
    score_table.style = "Light Grid Accent 1"
    hdr = score_table.rows[0].cells
    hdr[0].text, hdr[1].text = "维度", "得分"
    for name, val in _score_rows(report):
        row = score_table.add_row().cells
        row[0].text = name
        row[1].text = f"{val if val is not None else '-'}"

    # 问题统计
    doc.add_heading("问题统计", level=1)
    doc.add_paragraph(
        f"问题总数 {report.total_issues}　高优先级 {report.high_priority_issues}　"
        f"中优先级 {report.medium_priority_issues}　低优先级 {report.low_priority_issues}"
    )

    # 摘要 / 建议
    if report.summary:
        doc.add_heading("评审摘要", level=1)
        doc.add_paragraph(report.summary)
    if report.recommendations:
        doc.add_heading("改进建议", level=1)
        doc.add_paragraph(report.recommendations)

    # 模块评审结果
    module_results = list(report.module_results.select_related("module").all())
    if module_results:
        doc.add_heading("模块评审结果", level=1)
        for mr in module_results:
            mod_name = getattr(mr.module, "title", None) or "未命名模块"
            doc.add_heading(mod_name, level=2)
            try:
                rating = mr.get_module_rating_display()
            except Exception:
                rating = mr.module_rating or "-"
            doc.add_paragraph(f"评价：{rating}　问题数：{mr.issues_count}　严重度：{mr.severity_score}")
            for label, field in (("分析", mr.analysis_content), ("优点", mr.strengths),
                                 ("不足", mr.weaknesses), ("建议", mr.recommendations)):
                if field:
                    p = doc.add_paragraph()
                    p.add_run(f"{label}：").bold = True
                    p.add_run(field)

    # 问题清单
    issues = list(report.issues.select_related("module").all())
    if issues:
        doc.add_heading("问题清单", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(["优先级", "类型", "标题", "描述", "建议"]):
            table.rows[0].cells[i].text = h
        for it in issues:
            cells = table.add_row().cells
            cells[0].text = _issue_display(it, "priority")
            cells[1].text = _issue_display(it, "issue_type")
            cells[2].text = it.title or "-"
            cells[3].text = it.description or "-"
            cells[4].text = it.suggestion or "-"

    for p in doc.paragraphs:
        for run in p.runs:
            if not run.font.size:
                run.font.size = Pt(10.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
_CJK_FONT_REGISTERED = False
_CJK_FONT_NAME = "STSong-Light"


def _ensure_cjk_font():
    global _CJK_FONT_REGISTERED
    if _CJK_FONT_REGISTERED:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont(_CJK_FONT_NAME))
    _CJK_FONT_REGISTERED = True


def build_pdf(report: ReviewReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _ensure_cjk_font()

    styles = getSampleStyleSheet()
    base = ParagraphStyle("cjk", parent=styles["Normal"], fontName=_CJK_FONT_NAME,
                          fontSize=10, leading=15, alignment=TA_LEFT)
    h1 = ParagraphStyle("cjk_h1", parent=base, fontSize=16, leading=22, spaceBefore=6, spaceAfter=8)
    h2 = ParagraphStyle("cjk_h2", parent=base, fontSize=13, leading=18, spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("cjk_h3", parent=base, fontSize=11, leading=15, spaceBefore=6, spaceAfter=2)

    def esc(text) -> str:
        s = "" if text is None else str(text)
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    doc_title = getattr(getattr(report, "document", None), "title", None) or "需求评审报告"
    story.append(Paragraph(esc(f"{doc_title} — 需求评审报告"), h1))
    story.append(Paragraph(
        esc(f"评审人：{report.reviewer or 'AI需求评审助手'}　"
            f"评审时间：{report.review_date.strftime('%Y-%m-%d %H:%M') if report.review_date else '-'}　"
            f"总体评价：{_rating_text(report)}"),
        base,
    ))
    story.append(Spacer(1, 6))

    def table(data, col_widths=None):
        t = Table(data, colWidths=col_widths, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), _CJK_FONT_NAME),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f4ff")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    # 评分
    story.append(Paragraph("评分概览", h2))
    score_data = [["维度", "得分"]] + [[n, str(v if v is not None else "-")] for n, v in _score_rows(report)]
    story.append(table(score_data, col_widths=[60 * mm, 40 * mm]))
    story.append(Spacer(1, 6))

    story.append(Paragraph("问题统计", h2))
    story.append(Paragraph(
        esc(f"问题总数 {report.total_issues}　高优先级 {report.high_priority_issues}　"
            f"中优先级 {report.medium_priority_issues}　低优先级 {report.low_priority_issues}"),
        base,
    ))

    if report.summary:
        story.append(Paragraph("评审摘要", h2))
        story.append(Paragraph(esc(report.summary), base))
    if report.recommendations:
        story.append(Paragraph("改进建议", h2))
        story.append(Paragraph(esc(report.recommendations), base))

    module_results = list(report.module_results.select_related("module").all())
    if module_results:
        story.append(Paragraph("模块评审结果", h2))
        for mr in module_results:
            mod_name = getattr(mr.module, "title", None) or "未命名模块"
            story.append(Paragraph(esc(mod_name), h3))
            try:
                rating = mr.get_module_rating_display()
            except Exception:
                rating = mr.module_rating or "-"
            story.append(Paragraph(esc(f"评价：{rating}　问题数：{mr.issues_count}　严重度：{mr.severity_score}"), base))
            for label, field in (("分析", mr.analysis_content), ("优点", mr.strengths),
                                 ("不足", mr.weaknesses), ("建议", mr.recommendations)):
                if field:
                    story.append(Paragraph(esc(f"{label}：{field}"), base))

    issues = list(report.issues.select_related("module").all())
    if issues:
        story.append(Paragraph("问题清单", h2))
        issue_data = [["优先级", "类型", "标题", "描述", "建议"]]
        for it in issues:
            issue_data.append([
                Paragraph(esc(_issue_display(it, "priority")), base),
                Paragraph(esc(_issue_display(it, "issue_type")), base),
                Paragraph(esc(it.title or "-"), base),
                Paragraph(esc(it.description or "-"), base),
                Paragraph(esc(it.suggestion or "-"), base),
            ])
        story.append(table(issue_data, col_widths=[18 * mm, 22 * mm, 35 * mm, 50 * mm, 45 * mm]))

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    pdf.build(story)
    return buf.getvalue()


def build_export(report: ReviewReport, fmt: str):
    """返回 (filename, content_type, bytes)。fmt: 'pdf' | 'docx'/'word'。"""
    fmt = (fmt or "pdf").lower()
    stem = _report_filename_stem(report)
    if fmt in ("docx", "word"):
        return f"{stem}.docx", (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ), build_docx(report)
    return f"{stem}.pdf", "application/pdf", build_pdf(report)
